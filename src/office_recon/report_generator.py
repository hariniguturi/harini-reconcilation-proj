"""Write the formatted month-end branch workbook (openpyxl)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
TITLE_FONT = Font(bold=True, name="Calibri", size=14, color="1F4E79")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
STATUS_FILL = {
    "FULL_REVERSAL": PatternFill("solid", fgColor="C6EFCE"),
    "PART_REVERSAL": PatternFill("solid", fgColor="FFEB9C"),
    "PENDING": PatternFill("solid", fgColor="DDEBF7"),
    "OVERDUE": PatternFill("solid", fgColor="FFC7CE"),
    "OVER_REVERSED": PatternFill("solid", fgColor="F4B183"),
    "EXCEPTION": PatternFill("solid", fgColor="F4B183"),
    "CLOSED": PatternFill("solid", fgColor="C6EFCE"),
    "WITHIN_TAT": PatternFill("solid", fgColor="C6EFCE"),
}


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = 12
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value) + 3, 42))
        ws.column_dimensions[letter].width = width


def _write_sheet(ws, title: str, df: pd.DataFrame, status_cols: list[str]) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(df.columns), 1))

    start_row = 3
    if df.empty:
        ws.cell(start_row, 1, "No rows for this section.")
        return

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
            if r_idx == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
            elif df.columns[c_idx - 1] in status_cols:
                fill = STATUS_FILL.get(str(value))
                if fill is not None:
                    cell.fill = fill

    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"
    ws.freeze_panes = f"A{start_row + 1}"
    _autosize(ws)


def _summary_frame(reconciled: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "account_code",
        "account_type",
        "reference_no",
        "original_credit_ref",
        "amount",
        "reversal_status",
        "residual_balance",
        "ageing_days",
        "branch_code",
        "source",
        "linked_refs",
    ]
    present = [c for c in cols if c in reconciled.columns]
    out = reconciled[present].copy()
    out = out.rename(
        columns={
            "account_code": "Account Code",
            "account_type": "Account Type",
            "reference_no": "Reference No.",
            "original_credit_ref": "Original Credit Ref.",
            "amount": "Amount",
            "reversal_status": "Reversal Status",
            "residual_balance": "Residual Balance",
            "ageing_days": "Ageing (Days)",
            "branch_code": "Branch Code",
            "source": "Source",
            "linked_refs": "Linked Reversal Refs",
        }
    )
    return out


def _pending_frame(reconciled: pd.DataFrame) -> pd.DataFrame:
    mask = reconciled["reversal_status"].isin(["PENDING", "PART_REVERSAL"])
    cols = [
        "account_code",
        "reference_no",
        "residual_balance",
        "ageing_start_date",
        "ageing_days",
        "tat_status",
        "branch_code",
        "account_type",
    ]
    out = reconciled.loc[mask, [c for c in cols if c in reconciled.columns]].copy()
    return out.rename(
        columns={
            "account_code": "Account Code",
            "reference_no": "Reference No.",
            "residual_balance": "Pending Balance",
            "ageing_start_date": "Ageing Start Date",
            "ageing_days": "Days Pending",
            "tat_status": "TAT Status",
            "branch_code": "Branch",
            "account_type": "Account Type",
        }
    )


def _overdue_frame(reconciled: pd.DataFrame, branch: pd.DataFrame) -> pd.DataFrame:
    mask = reconciled["tat_status"] == "OVERDUE"
    frame = reconciled.loc[mask].copy()
    if not frame.empty and not branch.empty:
        frame = frame.merge(branch, on="branch_code", how="left")
    if "branch_name" not in frame.columns:
        frame["branch_name"] = ""
    if "email" not in frame.columns:
        frame["email"] = ""
    cols = [
        "account_code",
        "reference_no",
        "days_overdue",
        "branch_code",
        "branch_name",
        "email",
        "residual_balance",
        "ageing_days",
    ]
    out = frame[[c for c in cols if c in frame.columns]].copy()
    return out.rename(
        columns={
            "account_code": "Account Code",
            "reference_no": "Reference No.",
            "days_overdue": "Days Overdue",
            "branch_code": "Branch",
            "branch_name": "Assigned Owner",
            "email": "Email",
            "residual_balance": "Pending Balance",
            "ageing_days": "Ageing (Days)",
        }
    )


def processing_counts(reconciled: pd.DataFrame) -> dict[str, int]:
    return {
        "Total Entries": int(len(reconciled)),
        "Full Reversals": int((reconciled["reversal_status"] == "FULL_REVERSAL").sum()),
        "Part Reversals": int((reconciled["reversal_status"] == "PART_REVERSAL").sum()),
        "Pending": int((reconciled["reversal_status"] == "PENDING").sum()),
        "Overdue": int((reconciled["tat_status"] == "OVERDUE").sum()),
        "Over Reversed": int((reconciled["reversal_status"] == "OVER_REVERSED").sum()),
        "Exceptions": int((reconciled["tat_status"] == "EXCEPTION").sum()),
    }


def report_generator(
    reconciled: pd.DataFrame,
    branch: pd.DataFrame,
    unmatched_reversals: pd.DataFrame,
    counts: dict[str, int],
    output_path: Path,
    reporting_date: date,
    input_files: list[str],
    fifth_business_day: date | None = None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws_sum = wb.active
    ws_sum.title = "Reconciliation_Summary"
    _write_sheet(
        ws_sum,
        f"Office Account Reconciliation — {reporting_date.isoformat()}",
        _summary_frame(reconciled),
        ["Reversal Status"],
    )

    ws_pending = wb.create_sheet("Pending_PartReversal")
    _write_sheet(
        ws_pending,
        "Pending and part-reversal items carried forward",
        _pending_frame(reconciled),
        ["TAT Status"],
    )

    ws_overdue = wb.create_sheet("Overdue_Entries")
    _write_sheet(
        ws_overdue,
        "Entries breaching the 7-day TAT without a documented exception",
        _overdue_frame(reconciled, branch),
        [],
    )

    ws_proc = wb.create_sheet("Processing_Summary")
    ws_proc["A1"] = "Processing Summary"
    ws_proc["A1"].font = TITLE_FONT
    meta = [
        ("Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Reporting Date", reporting_date.isoformat()),
        ("Fifth Business Day Target", fifth_business_day.isoformat() if fifth_business_day else ""),
        ("Input Files", "; ".join(input_files)),
    ]
    for key, value in counts.items():
        meta.append((key, value))
    if not unmatched_reversals.empty:
        meta.append(("Unmatched Reversals", int(len(unmatched_reversals))))

    ws_proc["A3"] = "Metric"
    ws_proc["B3"] = "Value"
    ws_proc["A3"].fill = HEADER_FILL
    ws_proc["B3"].fill = HEADER_FILL
    ws_proc["A3"].font = HEADER_FONT
    ws_proc["B3"].font = HEADER_FONT
    for idx, (key, value) in enumerate(meta, start=4):
        ws_proc.cell(idx, 1, key).border = THIN
        ws_proc.cell(idx, 2, value).border = THIN
    _autosize(ws_proc)

    if not unmatched_reversals.empty:
        ws_unmatched = wb.create_sheet("Unmatched_Reversals")
        show_cols = [
            c
            for c in [
                "account_code",
                "reference_no",
                "original_credit_ref",
                "amount",
                "branch_code",
                "narration",
            ]
            if c in unmatched_reversals.columns
        ]
        _write_sheet(
            ws_unmatched,
            "Reversals that could not be linked to an original credit",
            unmatched_reversals[show_cols],
            [],
        )

    wb.save(output_path)
    return output_path
